import io

import docx
import pytest

from app.chunking import chunk_segments
from app.ingest import IngestError, parse_file, safe_filename

MB = 1024 * 1024


def test_parse_vietnamese_pdf_with_pages(fixtures):
    d = parse_file("vi_sample.pdf", (fixtures / "vi_sample.pdf").read_bytes(), 50 * MB)
    assert d.kind == "pdf" and d.pages == 3
    assert [s.page for s in d.segments] == [1, 2, 3]
    assert "Tây Nguyên" in d.segments[0].text
    assert "tưới nhỏ giọt" in d.segments[1].text.lower()
    chunks = chunk_segments(d.segments, 700, 100)
    assert any(c.page_end == 3 and "12,5 phần trăm" in c.text for c in chunks)


def test_parse_english_pdf(fixtures):
    d = parse_file("en.pdf", (fixtures / "en_sample.pdf").read_bytes(), 50 * MB)
    assert d.pages == 3 and "inverter" in d.segments[2].text


def test_parse_docx_generated_in_test():
    doc = docx.Document()
    doc.add_heading("Quy định nghỉ phép", 1)
    doc.add_paragraph("Nhân viên được nghỉ 12 ngày phép mỗi năm.")
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Thâm niên", "Ngày phép"
    t.cell(1, 0).text, t.cell(1, 1).text = "5 năm", "13 ngày"
    buf = io.BytesIO()
    doc.save(buf)
    d = parse_file("nghi-phep.docx", buf.getvalue(), 50 * MB)
    assert d.kind == "docx" and d.pages is None
    text = d.segments[0].text
    assert "12 ngày phép" in text and "Thâm niên | Ngày phép" in text and "5 năm | 13 ngày" in text


def test_parse_docx_fixture(fixtures):
    d = parse_file("vi_sample.docx", (fixtures / "vi_sample.docx").read_bytes(), 50 * MB)
    assert "Robusta" in d.segments[0].text


def test_parse_text_and_markdown(fixtures):
    d = parse_file("notes.md", (fixtures / "notes_en.md").read_bytes(), 50 * MB)
    assert d.kind == "txt" and "Solar" in d.segments[0].text
    d2 = parse_file("vi.txt", "Xin chào thế giới".encode("utf-8-sig"), 50 * MB)
    assert d2.segments[0].text == "Xin chào thế giới"


@pytest.mark.parametrize("name,code", [("scanned.pdf", "scanned"), ("corrupt.pdf", "corrupt")])
def test_bad_pdfs(fixtures, name, code):
    with pytest.raises(IngestError) as e:
        parse_file(name, (fixtures / name).read_bytes(), 50 * MB)
    assert e.value.code == code
    assert e.value.message  # có thông điệp tiếng Việt


def test_corrupt_docx_and_unsupported_and_too_large():
    with pytest.raises(IngestError) as e:
        parse_file("x.docx", b"PK\x03\x04 not really a zip", 50 * MB)
    assert e.value.code == "corrupt"
    with pytest.raises(IngestError) as e:
        parse_file("x.exe", b"MZ....", 50 * MB)
    assert e.value.code == "unsupported"
    with pytest.raises(IngestError) as e:
        parse_file("big.txt", b"a" * 2000, 1000)
    assert e.value.code == "too_large"


def test_safe_filename_strips_paths():
    assert safe_filename("../../etc/passwd") == "passwd"
    assert safe_filename("C:\\Users\\a\\Tài liệu.pdf") == "Tài liệu.pdf"
    assert safe_filename("") == "tai-lieu"
    assert "<" not in safe_filename("<script>.pdf")
    assert len(safe_filename("a" * 300 + ".pdf")) <= 120


def test_pdf_vietnamese_not_split_inside_syllables(fixtures):
    """pypdf hay chèn dấu cách giữa âm tiết ('tri ết h ọc'); PyMuPDF thì không."""
    import re
    d = parse_file("vi_sample.pdf", (fixtures / "vi_sample.pdf").read_bytes(), 50 * MB)
    text = "\n".join(s.text for s in d.segments)
    frag = re.compile(r"(?<!\w)(?:[bcdđghklmnpqrstvx]{1,3})\s+[ạảãàáâầấậẩẫăằắặẳẵẹẻẽèéêềếệểễịỉĩìíọỏõòóôồốộổỗơờớợởỡụủũùúưừứựửữỵỷỹỳý]\w*", re.I)
    assert not frag.findall(text)
